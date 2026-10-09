"""Builds data/<topic>/ folders from the Kaggle downloads plus a few generated files.

Re-runnable: it wipes and recreates each topic folder. Random choices use a fixed seed.
Ground-truth facts for the generated files are written to scripts/generated_facts.json
so the evaluation set can use them later.
"""
import json
import random
import shutil
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import pymupdf
from docx import Document
from PIL import Image, ImageFilter

DL = Path.home() / "Downloads"
ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
SEED = 42
rng = random.Random(SEED)
facts: dict = {}


def fresh(topic: str) -> Path:
    p = DATA / topic
    keep = (p / "topic.json").read_bytes() if (p / "topic.json").exists() else None
    if p.exists():
        shutil.rmtree(p)
    p.mkdir(parents=True)
    if keep is not None:  # hand-written topic description survives a rebuild
        (p / "topic.json").write_bytes(keep)
    return p


def strip_keys(obj, bad=("apikey",)):
    if isinstance(obj, dict):
        return {k: strip_keys(v, bad) for k, v in obj.items() if k not in bad}
    if isinstance(obj, list):
        return [strip_keys(v, bad) for v in obj]
    return obj


def bar_chart(series: pd.Series, title: str, path: Path, ylabel: str):
    fig, ax = plt.subplots(figsize=(7, 4))
    series.plot(kind="bar", ax=ax, color="#3b6ea5")
    ax.set_title(title)
    ax.set_ylabel(ylabel)
    for i, v in enumerate(series.values):
        ax.text(i, v, f"{v:g}", ha="center", va="bottom", fontsize=8)
    plt.xticks(rotation=30, ha="right")
    plt.tight_layout()
    fig.savefig(path, dpi=130)
    plt.close(fig)


def html_to_pdf(html: str, out: Path, resources: Path):
    story = pymupdf.Story(html=html, archive=pymupdf.Archive(str(resources)))
    writer = pymupdf.DocumentWriter(str(out))
    box = pymupdf.paper_rect("a4")
    where = box + (50, 50, -50, -50)
    more = 1
    while more:
        dev = writer.begin_page(box)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def scan_pdf(src_pdf: Path, out: Path):
    """Rasterise a text PDF into an image-only PDF (no text layer), like a scanner."""
    doc, new = pymupdf.open(src_pdf), pymupdf.open()
    tmp = out.with_suffix(".tmp.jpg")
    for page in doc:
        pix = page.get_pixmap(dpi=130)
        pix.save(str(tmp))
        img = Image.open(tmp).convert("L").rotate(0.8, expand=True, fillcolor=255)
        img = img.filter(ImageFilter.GaussianBlur(0.6))
        img.save(tmp, quality=55)
        w, h = img.size
        p = new.new_page(width=w * 72 / 130, height=h * 72 / 130)
        p.insert_image(p.rect, filename=str(tmp))
    new.save(out)
    tmp.unlink()


# ------------------------------------------------------------------ restaurants
def build_restaurants():
    d = fresh("restaurants")
    src = DL / "zomato.csv"
    z = pd.read_csv(src / "zomato.csv", encoding="latin-1")
    z.to_csv(d / "zomato_restaurants.csv", index=False)  # re-saved as UTF-8
    shutil.copy(src / next(f for f in (p.name for p in src.iterdir()) if f.endswith(".xlsx")),
                d / "country_codes.xlsx")
    j = json.load(open(src / "file1.json", encoding="utf-8"))
    j = strip_keys(j)
    merged = [r for page in j for r in page.get("restaurants", [])][:100]
    j = {"restaurants": merged, "results_shown": len(merged)}
    json.dump(j, open(d / "delhi_restaurants_sample.json", "w", encoding="utf-8"), indent=1)

    top = z[z["City"].isin(z["City"].value_counts().head(6).index)]
    avg = top.groupby("City")["Aggregate rating"].mean().round(2).sort_values()
    bar_chart(avg, "Average restaurant rating by city", d / "avg_rating_by_city_chart.png", "Avg rating")
    facts["restaurants"] = {"chart_avg_rating_by_city": avg.to_dict(),
                            "total_rows": len(z), "mean_rating_all": round(z["Aggregate rating"].mean(), 3)}

    html = """<html><body><h1>Zomato Dining Guide</h1>
    <h2>Online delivery</h2><p>Many restaurants in the dataset offer online delivery, which customers
    can use to order food from home. The table-booking option is far less common.</p>
    <h2>Price range</h2><p>Price range is a number from 1 (cheapest) to 4 (most expensive).</p>
    </body></html>"""
    (d / "dining_guide.html").write_text(html, encoding="utf-8")


# ------------------------------------------------------------------ fashion
def build_fashion():
    d = fresh("fashion_products")
    base = DL / "fashion product"
    styles = pd.read_csv(base / "myntradataset" / "styles.csv", on_bad_lines="skip")
    img_dir = base / "images"
    have = {int(p.stem) for p in img_dir.glob("*.jpg")}
    styles = styles[styles["id"].isin(have)].dropna(subset=["productDisplayName"])
    # 30 images, spread over article types so image questions are varied
    picks = (styles.groupby("articleType", group_keys=False)
             .apply(lambda g: g.sample(min(len(g), 2), random_state=SEED))
             .sample(30, random_state=SEED))
    rest = styles.drop(picks.index).sample(470, random_state=SEED)
    sample = pd.concat([picks, rest]).sample(frac=1, random_state=SEED)
    sample.to_csv(d / "styles_catalog.csv", index=False)
    (d / "images").mkdir()
    for i in picks["id"]:
        shutil.copy(img_dir / f"{i}.jpg", d / "images" / f"{i}.jpg")

    fk = pd.read_excel(DL / "ecommerce product" / "output.xlsx")
    fk = fk.drop(columns=[c for c in ("Unnamed: 0", "_id", "images", "url", "crawled_at", "pid") if c in fk])
    fk.sample(300, random_state=SEED).to_excel(d / "flipkart_products.xlsx", index=False)

    counts = sample["articleType"].value_counts().head(8)
    bar_chart(counts, "Products per article type (catalog sample)", d / "articletype_counts_chart.png", "Products")
    facts["fashion_products"] = {"catalog_rows": len(sample), "image_ids": picks["id"].tolist(),
                                 "chart_counts": counts.to_dict()}


# ------------------------------------------------------------------ airline tweets
def build_airline():
    d = fresh("airline_tweets")
    t = pd.read_csv(DL / "tweets.csv" / "Tweets.csv")
    t.to_csv(d / "airline_tweets.csv", index=False)
    neg = t[(t["airline_sentiment"] == "negative") & t["negativereason"].notna()]
    n = 0
    for airline, g in neg.groupby("airline"):
        for reason, gg in g.groupby("negativereason"):
            if len(gg) < 8 or n >= 20:
                continue
            lines = [f"# {airline}: customers complaining about '{reason}'", ""]
            lines += [f"- {r.tweet_created[:10]} @{r['name']}: {r.text}" for _, r in gg.head(12).iterrows()]
            fn = f"{airline.replace(' ', '_').lower()}_{reason.replace(' ', '_').replace('/', '').lower()}.txt"
            (d / fn).write_text("\n".join(lines), encoding="utf-8")
            n += 1

    doc = Document()
    doc.add_heading("Airline Customer Care Policy", 0)
    doc.add_heading("Delays", 1)
    doc.add_paragraph("Passengers delayed more than 3 hours are entitled to a meal voucher.")
    doc.add_heading("Lost luggage", 1)
    doc.add_paragraph("Lost luggage must be reported within 24 hours. Compensation is capped at 1500 dollars.")
    doc.save(d / "customer_care_policy.docx")
    facts["airline_tweets"] = {"rows": len(t), "negative": int((t.airline_sentiment == "negative").sum()),
                               "by_airline": t["airline"].value_counts().to_dict(), "thread_files": n,
                               "docx": "delay>3h -> meal voucher; lost luggage within 24h; cap 1500 dollars"}


# ------------------------------------------------------------------ song lyrics
def build_songs():
    d = fresh("song_lyrics")
    s = pd.read_csv(DL / "spotify.csv" / "Spotify Million Song Dataset_exported.csv")
    s = s[s["text"].str.len().between(600, 2500)].dropna(subset=["song", "artist"])
    pick = s.sample(25, random_state=SEED)
    for _, r in pick.iterrows():
        safe = "".join(c for c in f"{r.artist} - {r.song}" if c.isalnum() or c in " -_")[:80].strip()
        (d / f"{safe}.txt").write_text(f"{r.song}\nby {r.artist}\n\n{r.text.strip()}", encoding="utf-8")
    pick[["artist", "song"]].assign(lyric_chars=pick["text"].str.len()).to_csv(d / "songs_index.csv", index=False)
    facts["song_lyrics"] = {"songs": len(pick)}


# ------------------------------------------------------------------ wildlife reports
def build_wildlife():
    d = fresh("wildlife_reports")
    p = pd.read_parquet(DL / "archive (3)" / "data.parquet")
    c = p["original_completion"].fillna("")
    # The parquet mixes unrelated topics (finance, farming, ...); keep only snake/wildlife reports.
    p = p[c.str.contains("Displacement Risk") & c.str.contains(r"\|") & c.str.len().between(2500, 9000)]
    assert len(p) >= 15, f"only {len(p)} snake reports match"
    for i, (_, r) in enumerate(p.sample(15, random_state=SEED).iterrows(), 1):
        (d / f"risk_assessment_{i:02d}.md").write_text(r["original_completion"], encoding="utf-8")

    # generated PDF: text + ruled table + embedded chart
    rs = random.Random(SEED)
    cities = ["Nairobi", "Mombasa", "Kisumu", "Nakuru", "Eldoret"]
    rows = [(c, rs.randint(20, 90), rs.randint(5, 30)) for c in cities]
    res = ROOT / "scripts" / "_tmp"
    res.mkdir(exist_ok=True)
    bar_chart(pd.Series({c: s for c, s, _ in rows}), "Snake sightings by city, 2025", res / "sightings.png", "Sightings")
    tr = "".join(f"<tr><td>{c}</td><td>{s}</td><td>{r}</td></tr>" for c, s, r in rows)
    html = f"""<html><head><style>
    table {{ border-collapse: collapse; }}
    th, td {{ border: 1px solid #000000; padding: 5px; }}
    th {{ background-color: #dddddd; }}
    </style></head><body><h1>Urban Wildlife Survey 2025</h1>
    <p>This survey records snake sightings reported by residents in five Kenyan cities during 2025.
    Most sightings were harmless species. The relocation programme moved snakes to protected land.</p>
    <h2>Sightings and relocations</h2>
    <table border="1" cellpadding="4"><tr><th>City</th><th>Sightings</th><th>Relocations</th></tr>{tr}</table>
    <h2>Chart</h2><img src="sightings.png" width="420"/>
    <p>Figure 1: Snake sightings by city in 2025.</p></body></html>"""
    html_to_pdf(html, d / "urban_wildlife_survey_2025.pdf", res)

    memo = """<html><body><h1>Field Memo 14 March 2025</h1>
    <p>On 14 March 2025 the Karen team relocated 12 green snakes to the Ngong forest reserve.
    The team lead noted that two residents were bitten by non-venomous snakes and treated without incident.
    Next inspection is planned for 2 April 2025 with budget approved at 4800 dollars.</p></body></html>"""
    html_to_pdf(memo, res / "memo.pdf", res)
    scan_pdf(res / "memo.pdf", d / "field_memo_scanned.pdf")
    shutil.rmtree(res)
    facts["wildlife_reports"] = {"survey_table": rows, "total_sightings": sum(s for _, s, _ in rows),
                                 "total_relocations": sum(r for _, _, r in rows),
                                 "scanned_memo": "12 green snakes, Ngong forest reserve, next inspection 2 April 2025, budget 4800 dollars"}


# ------------------------------------------------------------------ receipts
def build_receipts():
    d = fresh("receipts")
    base = DL / "receipt" / "SROIE2019" / "train"
    ids = sorted(p.stem for p in (base / "img").glob("*.jpg"))
    pick = rng.sample(ids, 31)
    rows = []
    for i in pick[:25]:
        shutil.copy(base / "img" / f"{i}.jpg", d / f"{i}.jpg")
        e = json.load(open(base / "entities" / f"{i}.txt", encoding="utf-8"))
        rows.append({"receipt_file": f"{i}.jpg", **e})
    pd.DataFrame(rows).to_csv(d / "receipts_ground_truth.csv", index=False)

    scans = {}
    for n, chunk in enumerate((pick[25:28], pick[28:31]), 1):
        doc = pymupdf.open()
        for i in chunk:
            img = pymupdf.open(base / "img" / f"{i}.jpg")
            pg = doc.new_page(width=img[0].rect.width, height=img[0].rect.height)
            pg.insert_image(pg.rect, filename=str(base / "img" / f"{i}.jpg"))
            scans[f"receipts_scan_batch{n}.pdf p.{len(doc)}"] = json.load(
                open(base / "entities" / f"{i}.txt", encoding="utf-8"))
        doc.save(d / f"receipts_scan_batch{n}.pdf")
    totals = [float(str(r["total"]).replace(",", "")) for r in rows if str(r["total"]).replace(",", "").replace(".", "").isdigit()]
    facts["receipts"] = {"csv_rows": len(rows), "sum_total": round(sum(totals), 2), "scan_pdfs": scans}


if __name__ == "__main__":
    DATA.mkdir(exist_ok=True)
    for fn in (build_restaurants, build_fashion, build_airline, build_songs, build_wildlife, build_receipts):
        print("building", fn.__name__)
        fn()
    json.dump(facts, open(ROOT / "scripts" / "generated_facts.json", "w"), indent=1, default=str)
    print("done")
