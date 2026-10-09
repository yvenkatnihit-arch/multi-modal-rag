from docx import Document

from src.ingestion.file_router import RoutedFile
from src.ingestion.parsers.common import clean_text
from src.ingestion.parsers.text_parser import parse_text


def parse(tmp_path, name, content: str | bytes):
    p = tmp_path / name
    p.write_bytes(content if isinstance(content, bytes) else content.encode("utf-8"))
    return parse_text(RoutedFile(p, name, "text"), "t")


def test_clean_text_fixes_ligatures_and_whitespace():
    assert clean_text("ﬁve snakes \x07 here\r\n\r\n\r\n\r\nend") == "five snakes here\n\nend"


def test_markdown_headings_and_paths(tmp_path):
    recs = parse(tmp_path, "r.md", "# Report\nintro\n\n## Risk\nhigh\n\n### Detail\nmore\n\n## Fix\nact\n")
    assert [(r.metadata.heading, r.text) for r in recs] == [
        ("Report", "intro"), ("Report > Risk", "high"),
        ("Report > Risk > Detail", "more"), ("Report > Fix", "act"),
    ]


def test_hash_inside_code_fence_is_not_a_heading(tmp_path):
    recs = parse(tmp_path, "c.md", "# A\ntext\n```\n# not a heading\n```\n")
    assert len(recs) == 1 and "# not a heading" in recs[0].text


def test_heading_only_sections_are_dropped_but_keep_path(tmp_path):
    recs = parse(tmp_path, "h.md", "# A\n## B\nbody\n")
    assert [(r.metadata.heading, r.text) for r in recs] == [("A > B", "body")]


def test_separator_only_sections_are_dropped(tmp_path):
    recs = parse(tmp_path, "s.md", "# A\nreal text\n\n## B\n---\n\n## C\n***\nmore text\n")
    assert [(r.metadata.heading, r.text) for r in recs] == [("A", "real text"), ("A > C", "***\nmore text")]


def test_plain_txt_without_headings_is_one_record(tmp_path):
    recs = parse(tmp_path, "song.txt", "Title\nby Someone\n\nla la la\n\nla la")
    assert len(recs) == 1 and recs[0].metadata.heading is None and "la la" in recs[0].text


def test_html_strips_scripts_and_keeps_structure(tmp_path):
    html = ("<html><head><style>x{}</style></head><body><h1>Guide</h1><script>bad()</script>"
            "<p>Intro</p><h2>Price</h2><p>Range 1-4</p>"
            "<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>")
    recs = parse(tmp_path, "g.html", html)
    assert [r.metadata.heading for r in recs] == ["Guide", "Guide > Price"]
    assert "bad()" not in " ".join(r.text for r in recs)
    assert "A | B" in recs[1].text and "1 | 2" in recs[1].text


def test_docx_headings_and_tables(tmp_path):
    d = Document()
    d.add_heading("Policy", 0)
    d.add_heading("Delays", 1)
    d.add_paragraph("Over 3 hours: voucher.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Cap", "Luggage", "1500"
    p = tmp_path / "p.docx"
    d.save(p)
    recs = parse_text(RoutedFile(p, "p.docx", "text"), "t")
    assert recs[0].metadata.heading == "Policy > Delays"
    assert "voucher" in recs[0].text and "Luggage | 1500" in recs[0].text


def test_cp1252_file_is_decoded(tmp_path):
    recs = parse(tmp_path, "w.txt", b"caf\xe9 \x93quoted\x94")   # raw Windows-1252 bytes
    assert "café" in recs[0].text and "“quoted”" in recs[0].text


def test_whitespace_only_file_gives_no_records(tmp_path):
    assert parse(tmp_path, "e.txt", "  \n\n \n") == []
